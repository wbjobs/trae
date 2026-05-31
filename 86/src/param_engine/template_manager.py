import csv
import json
import os
from typing import Dict, List, Optional, Any, Tuple
from dataclasses import dataclass, field
from pathlib import Path
from enum import Enum
import uuid
from src.core.models import DeviceConfig, TagPoint, PointType, DataType


class TemplateFormat(Enum):
    CSV = "csv"
    JSON = "json"
    EXCEL = "excel"


@dataclass
class TagImportResult:
    tag_name: str
    success: bool
    message: str = ""
    old_value: Any = None
    new_value: Any = None


@dataclass
class BatchImportResult:
    total: int = 0
    success: int = 0
    failed: int = 0
    results: List[TagImportResult] = field(default_factory=list)
    errors: List[str] = field(default_factory=list)


class ParameterTemplateManager:
    REQUIRED_COLUMNS = ['name', 'address', 'point_type', 'data_type']
    OPTIONAL_COLUMNS = [
        'description', 'unit', 'min_value', 'max_value',
        'default_value', 'coefficient', 'offset', 'group',
        'alarm_low', 'alarm_high', 'read_only'
    ]

    def __init__(self, project_manager):
        self._project_manager = project_manager

    def generate_template(self, format: TemplateFormat, output_path: str, include_sample: bool = True) -> bool:
        try:
            if format == TemplateFormat.CSV:
                return self._generate_csv_template(output_path, include_sample)
            elif format == TemplateFormat.JSON:
                return self._generate_json_template(output_path, include_sample)
            elif format == TemplateFormat.EXCEL:
                return self._generate_excel_template(output_path, include_sample)
            return False
        except Exception as e:
            print(f"Generate template error: {e}")
            return False

    def _generate_csv_template(self, output_path: str, include_sample: bool) -> bool:
        columns = self.REQUIRED_COLUMNS + self.OPTIONAL_COLUMNS
        with open(output_path, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            writer.writerow(columns)
            if include_sample:
                writer.writerow([
                    'AI_TEMP_001', '40001', 'AI', 'float32',
                    '温度传感器1', '℃', '-50', '200',
                    '25', '1.0', '0.0', '温度',
                    '0', '100', 'True'
                ])
                writer.writerow([
                    'DO_VALVE_001', '0', 'DO', 'bool',
                    '电磁阀1', '', '0', '1',
                    '0', '1.0', '0.0', '控制',
                    '', '', 'False'
                ])
        return True

    def _generate_json_template(self, output_path: str, include_sample: bool) -> bool:
        template = {
            "version": "1.0",
            "description": "批量点位参数导入模板",
            "tags": []
        }
        if include_sample:
            template["tags"] = [
                {
                    "name": "AI_TEMP_001",
                    "address": "40001",
                    "point_type": "AI",
                    "data_type": "float32",
                    "description": "温度传感器1",
                    "unit": "℃",
                    "min_value": -50,
                    "max_value": 200,
                    "default_value": 25,
                    "coefficient": 1.0,
                    "offset": 0.0,
                    "group": "温度",
                    "alarm_low": 0,
                    "alarm_high": 100,
                    "read_only": True
                },
                {
                    "name": "DO_VALVE_001",
                    "address": "0",
                    "point_type": "DO",
                    "data_type": "bool",
                    "description": "电磁阀1",
                    "unit": "",
                    "min_value": 0,
                    "max_value": 1,
                    "default_value": 0,
                    "coefficient": 1.0,
                    "offset": 0.0,
                    "group": "控制",
                    "read_only": False
                }
            ]
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(template, f, indent=4, ensure_ascii=False)
        return True

    def _generate_excel_template(self, output_path: str, include_sample: bool) -> bool:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            print("openpyxl not installed, falling back to CSV")
            return self._generate_csv_template(output_path.replace('.xlsx', '.csv'), include_sample)
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "点位参数"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        required_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        center_align = Alignment(horizontal='center', vertical='center')

        columns = self.REQUIRED_COLUMNS + self.OPTIONAL_COLUMNS
        for col, header in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            if header in self.REQUIRED_COLUMNS:
                cell.fill = required_fill

        if include_sample:
            sample_data = [
                ['AI_TEMP_001', '40001', 'AI', 'float32',
                 '温度传感器1', '℃', -50, 200, 25, 1.0, 0.0, '温度', 0, 100, True],
                ['DO_VALVE_001', '0', 'DO', 'bool',
                 '电磁阀1', '', 0, 1, 0, 1.0, 0.0, '控制', None, None, False]
            ]
            for row_idx, data in enumerate(sample_data, 2):
                for col_idx, value in enumerate(data, 1):
                    ws.cell(row=row_idx, column=col_idx, value=value)

        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[column].width = max_length + 2

        wb.save(output_path)
        return True

    def import_template(self, file_path: str, target_device_id: str, 
                        update_existing: bool = True, create_missing: bool = True) -> BatchImportResult:
        result = BatchImportResult()
        project = self._project_manager.current_project
        if not project:
            result.errors.append("没有打开的工程")
            return result

        device = project.devices.get(target_device_id)
        if not device:
            result.errors.append(f"设备不存在: {target_device_id}")
            return result

        ext = Path(file_path).suffix.lower()
        try:
            if ext == '.csv':
                tags_data = self._parse_csv_template(file_path)
            elif ext == '.json':
                tags_data = self._parse_json_template(file_path)
            elif ext in ['.xlsx', '.xls']:
                tags_data = self._parse_excel_template(file_path)
            else:
                result.errors.append(f"不支持的文件格式: {ext}")
                return result
        except Exception as e:
            result.errors.append(f"解析模板失败: {e}")
            return result

        result.total = len(tags_data)
        for tag_data in tags_data:
            import_result = self._import_single_tag(tag_data, device, update_existing, create_missing)
            result.results.append(import_result)
            if import_result.success:
                result.success += 1
            else:
                result.failed += 1

        return result

    def _parse_csv_template(self, file_path: str) -> List[Dict[str, Any]]:
        tags_data = []
        encoding = 'utf-8-sig' if os.name == 'nt' else 'utf-8'
        try:
            with open(file_path, 'r', encoding=encoding) as f:
                reader = csv.DictReader(f)
                for row in reader:
                    if any(row.values()):
                        tags_data.append(self._clean_row_data(row))
        except UnicodeDecodeError:
            with open(file_path, 'r', encoding='gbk') as f:
                reader = csv.DictReader(f)
                for row in reader:
                    if any(row.values()):
                        tags_data.append(self._clean_row_data(row))
        return tags_data

    def _parse_json_template(self, file_path: str) -> List[Dict[str, Any]]:
        with open(file_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        return data.get("tags", [])

    def _parse_excel_template(self, file_path: str) -> List[Dict[str, Any]]:
        try:
            import openpyxl
        except ImportError:
            raise ImportError("需要安装 openpyxl 库来处理Excel文件")
        
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
        
        headers = []
        for cell in ws[1]:
            headers.append(cell.value)
        
        tags_data = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if any(v is not None for v in row):
                row_dict = {}
                for header, value in zip(headers, row):
                    if header:
                        row_dict[header] = value
                tags_data.append(self._clean_row_data(row_dict))
        
        return tags_data

    def _clean_row_data(self, row: Dict[str, Any]) -> Dict[str, Any]:
        cleaned = {}
        for key, value in row.items():
            if value is None:
                continue
            key_lower = key.lower()
            if key_lower in ['min_value', 'max_value', 'default_value', 'coefficient', 'offset', 'alarm_low', 'alarm_high']:
                try:
                    cleaned[key_lower] = float(value) if value != '' else None
                except (ValueError, TypeError):
                    cleaned[key_lower] = None
            elif key_lower == 'read_only':
                if isinstance(value, bool):
                    cleaned[key_lower] = value
                else:
                    cleaned[key_lower] = str(value).lower() in ['true', '1', 'yes', '是']
            else:
                cleaned[key_lower] = str(value).strip() if value else ''
        return cleaned

    def _import_single_tag(self, tag_data: Dict[str, Any], device: DeviceConfig,
                           update_existing: bool, create_missing: bool) -> TagImportResult:
        tag_name = tag_data.get('name', '').strip()
        if not tag_name:
            return TagImportResult(tag_name="", success=False, message="点位名称为空")

        for col in self.REQUIRED_COLUMNS:
            if col not in tag_data or not tag_data[col]:
                return TagImportResult(
                    tag_name=tag_name,
                    success=False,
                    message=f"缺少必填列: {col}"
                )

        try:
            point_type = PointType(tag_data['point_type'])
            data_type = DataType(tag_data['data_type'])
        except ValueError as e:
            return TagImportResult(
                tag_name=tag_name,
                success=False,
                message=f"无效的类型: {e}"
            )

        existing_tag = device.tags.get(tag_name)
        old_value = None

        if existing_tag:
            if not update_existing:
                return TagImportResult(
                    tag_name=tag_name,
                    success=False,
                    message="点位已存在且不允许更新"
                )
            old_value = {
                'address': existing_tag.address,
                'coefficient': existing_tag.coefficient,
                'offset': existing_tag.offset,
                'min_value': existing_tag.min_value,
                'max_value': existing_tag.max_value
            }
            tag = existing_tag
        else:
            if not create_missing:
                return TagImportResult(
                    tag_name=tag_name,
                    success=False,
                    message="点位不存在且不允许创建"
                )
            tag = TagPoint(
                name=tag_name,
                address=tag_data['address'],
                point_type=point_type,
                data_type=data_type
            )
            device.tags[tag_name] = tag

        tag.address = tag_data.get('address', tag.address)
        tag.point_type = point_type
        tag.data_type = data_type
        tag.description = tag_data.get('description', tag.description)
        tag.unit = tag_data.get('unit', tag.unit)
        tag.min_value = tag_data.get('min_value', tag.min_value) or 0.0
        tag.max_value = tag_data.get('max_value', tag.max_value) or 100.0
        tag.default_value = tag_data.get('default_value', tag.default_value)
        tag.coefficient = tag_data.get('coefficient', tag.coefficient) or 1.0
        tag.offset = tag_data.get('offset', tag.offset) or 0.0
        tag.group = tag_data.get('group', tag.group)
        tag.alarm_low = tag_data.get('alarm_low', tag.alarm_low)
        tag.alarm_high = tag_data.get('alarm_high', tag.alarm_high)
        tag.read_only = tag_data.get('read_only', tag.read_only)

        new_value = {
            'address': tag.address,
            'coefficient': tag.coefficient,
            'offset': tag.offset,
            'min_value': tag.min_value,
            'max_value': tag.max_value
        }

        return TagImportResult(
            tag_name=tag_name,
            success=True,
            message="导入成功",
            old_value=old_value,
            new_value=new_value
        )

    def batch_apply_parameters(self, parameters: Dict[str, Dict[str, Any]], 
                               target_device_id: str = None) -> BatchImportResult:
        result = BatchImportResult()
        project = self._project_manager.current_project
        if not project:
            result.errors.append("没有打开的工程")
            return result

        devices = [project.devices[target_device_id]] if target_device_id else list(project.devices.values())
        
        for device in devices:
            for tag_name, params in parameters.items():
                if tag_name in device.tags:
                    tag = device.tags[tag_name]
                    old_value = {k: getattr(tag, k) for k in params if hasattr(tag, k)}
                    for key, value in params.items():
                        if hasattr(tag, key):
                            setattr(tag, key, value)
                    result.results.append(TagImportResult(
                        tag_name=tag_name,
                        success=True,
                        message="参数更新成功",
                        old_value=old_value,
                        new_value=params
                    ))
                    result.success += 1
                else:
                    result.results.append(TagImportResult(
                        tag_name=tag_name,
                        success=False,
                        message="点位不存在"
                    ))
                    result.failed += 1
                result.total += 1

        return result

    def export_parameters(self, output_path: str, device_id: str = None, 
                          group_filter: str = None) -> bool:
        project = self._project_manager.current_project
        if not project:
            return False

        tags_data = []
        devices = [project.devices[device_id]] if device_id else list(project.devices.values())
        
        for device in devices:
            for tag in device.tags.values():
                if group_filter and tag.group != group_filter:
                    continue
                tags_data.append({
                    'device_id': device.device_id,
                    'device_name': device.device_name,
                    'name': tag.name,
                    'address': tag.address,
                    'point_type': tag.point_type.value,
                    'data_type': tag.data_type.value,
                    'description': tag.description,
                    'unit': tag.unit,
                    'min_value': tag.min_value,
                    'max_value': tag.max_value,
                    'current_value': tag.current_value,
                    'coefficient': tag.coefficient,
                    'offset': tag.offset,
                    'group': tag.group,
                    'read_only': tag.read_only
                })

        ext = Path(output_path).suffix.lower()
        try:
            if ext == '.csv':
                self._export_to_csv(output_path, tags_data)
            elif ext == '.json':
                self._export_to_json(output_path, tags_data)
            elif ext in ['.xlsx', '.xls']:
                self._export_to_excel(output_path, tags_data)
            else:
                return False
            return True
        except Exception as e:
            print(f"Export error: {e}")
            return False

    def _export_to_csv(self, output_path: str, tags_data: List[Dict]):
        if not tags_data:
            return
        columns = list(tags_data[0].keys())
        encoding = 'utf-8-sig' if os.name == 'nt' else 'utf-8'
        with open(output_path, 'w', encoding=encoding, newline='') as f:
            writer = csv.DictWriter(f, fieldnames=columns)
            writer.writeheader()
            writer.writerows(tags_data)

    def _export_to_json(self, output_path: str, tags_data: List[Dict]):
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump({"version": "1.0", "tags": tags_data}, f, indent=4, ensure_ascii=False)

    def _export_to_excel(self, output_path: str, tags_data: List[Dict]):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            self._export_to_csv(output_path.replace('.xlsx', '.csv'), tags_data)
            return

        if not tags_data:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "点位参数"

        columns = list(tags_data[0].keys())
        for col, header in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for row_idx, tag_data in enumerate(tags_data, 2):
            for col_idx, key in enumerate(columns, 1):
                ws.cell(row=row_idx, column=col_idx, value=tag_data.get(key))

        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[column].width = max_length + 2

        wb.save(output_path)
