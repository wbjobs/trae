import io
import os
from typing import Tuple, List, Dict, Any
from datetime import datetime
from docx import Document
from docx.oxml.ns import qn
from PIL import Image
from .base_parser import BaseParser
from ..schemas.document import TableData, ImageData


class WordParser(BaseParser):
    def __init__(self, file_path: str, original_filename: str):
        super().__init__(file_path, original_filename)
        self.file_type = "word"

    def parse(self) -> Tuple[str, List[TableData], List[ImageData], Dict[str, Any]]:
        text_content_parts: List[str] = []
        tables: List[TableData] = []
        images: List[ImageData] = []
        metadata: Dict[str, Any] = {}

        doc = Document(self.file_path)

        core_props = doc.core_properties
        if core_props:
            metadata["title"] = core_props.title
            metadata["author"] = core_props.author
            if core_props.created:
                metadata["creation_date"] = core_props.created

        for para in doc.paragraphs:
            if para.text.strip():
                text_content_parts.append(para.text)

        for table_idx, table in enumerate(doc.tables):
            headers: List[str] = []
            rows: List[List[str]] = []

            for row_idx, row in enumerate(table.rows):
                cells = [cell.text.strip() for cell in row.cells]
                if row_idx == 0:
                    headers = cells
                else:
                    rows.append(cells)

            if headers or rows:
                csv_content = self._table_to_csv(headers, rows)
                tables.append(TableData(
                    headers=headers,
                    rows=rows,
                    csv_content=csv_content
                ))

        images = self._extract_images(doc)

        metadata["file_type"] = self.file_type
        text_content = "\n\n".join(text_content_parts)

        return text_content, tables, images, metadata

    def _table_to_csv(self, headers: List[str], rows: List[List[str]]) -> str:
        csv_lines = []
        if headers:
            csv_lines.append(",".join(headers))
        for row in rows:
            csv_lines.append(",".join([f'"{cell}"' if "," in cell else cell for cell in row]))
        return "\n".join(csv_lines)

    def _extract_images(self, doc: Document) -> List[ImageData]:
        extracted_images: List[ImageData] = []

        try:
            for rel in doc.part.rels.values():
                if "image" in rel.target_ref:
                    try:
                        image_part = rel.target_part
                        image_bytes = image_part.blob
                        ext = image_part.content_type.split("/")[-1]
                        if ext == "jpeg":
                            ext = "jpg"

                        image_data = self.save_image(image_bytes, ext)

                        try:
                            with Image.open(io.BytesIO(image_bytes)) as pil_img:
                                image_data.width, image_data.height = pil_img.size
                        except Exception:
                            pass

                        extracted_images.append(image_data)
                    except Exception:
                        continue
        except Exception:
            pass

        return extracted_images
