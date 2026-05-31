import os
import re
import logging
from typing import Optional, Dict, Any, List, Tuple
from dataclasses import dataclass
from pathlib import Path

logger = logging.getLogger(__name__)


class EncodingDetector:
    COMMON_ENCODINGS = [
        "utf-8", "utf-8-sig", "gbk", "gb2312", "gb18030", "big5",
        "utf-16", "utf-16-le", "utf-16-be",
        "ascii", "latin-1", "cp1252", "cp936", "shift_jis", "euc-jp", "euc-kr"
    ]

    @staticmethod
    def detect_and_read(file_path: str, max_lines: int = 100) -> Tuple[str, str]:
        raw_bytes = b""
        try:
            with open(file_path, "rb") as f:
                raw_bytes = f.read()
        except Exception as e:
            logger.error(f"Failed to read file bytes: {e}")
            return "", "unknown"

        if not raw_bytes:
            return "", "empty"

        detected_encoding = EncodingDetector._detect_encoding(raw_bytes)
        if detected_encoding:
            try:
                text = raw_bytes.decode(detected_encoding)
                if EncodingDetector._is_valid_text(text):
                    return text, detected_encoding
            except (UnicodeDecodeError, LookupError):
                pass

        for encoding in EncodingDetector.COMMON_ENCODINGS:
            try:
                text = raw_bytes.decode(encoding, errors="strict")
                if EncodingDetector._is_valid_text(text):
                    logger.info(f"Successfully decoded with encoding: {encoding}")
                    return text, encoding
            except (UnicodeDecodeError, LookupError):
                continue

        logger.warning("Falling back to UTF-8 with replacement")
        text = raw_bytes.decode("utf-8", errors="replace")
        text = EncodingDetector._clean_garbage_chars(text)
        return text, "utf-8-replaced"

    @staticmethod
    def _detect_encoding(raw_bytes: bytes) -> Optional[str]:
        try:
            import chardet
            result = chardet.detect(raw_bytes[:4096])
            if result["confidence"] > 0.7:
                return result["encoding"]
        except ImportError:
            pass

        if raw_bytes.startswith(b"\xef\xbb\xbf"):
            return "utf-8-sig"
        if raw_bytes.startswith(b"\xff\xfe\x00\x00") or raw_bytes.startswith(b"\x00\x00\xfe\xff"):
            return "utf-32"
        if raw_bytes.startswith(b"\xff\xfe") or raw_bytes.startswith(b"\xfe\xff"):
            return "utf-16"

        return None

    @staticmethod
    def _is_valid_text(text: str) -> bool:
        if not text:
            return False
        printable_ratio = sum(1 for c in text if c.isprintable() or c in "\n\r\t") / len(text)
        if printable_ratio < 0.9:
            return False

        control_chars = sum(1 for c in text if ord(c) < 32 and c not in "\n\r\t")
        if control_chars > len(text) * 0.05:
            return False

        return True

    @staticmethod
    def _clean_garbage_chars(text: str) -> str:
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
        text = re.sub(r"\ufffd+", " ", text)
        text = re.sub(r"[ \t]+", " ", text)
        return text


@dataclass
class DocumentContent:
    text: str
    metadata: Dict[str, Any]
    tables: List[List[List[str]]]
    images: List[bytes]
    file_path: str
    file_type: str


class BaseFormatAdapter:
    def __init__(self, config: Any = None):
        self.config = config

    def can_handle(self, file_path: str) -> bool:
        raise NotImplementedError

    def extract(self, file_path: str) -> DocumentContent:
        raise NotImplementedError


class TextAdapter(BaseFormatAdapter):
    def can_handle(self, file_path: str) -> bool:
        ext = Path(file_path).suffix.lower()
        return ext in [".txt", ".md", ".log", ".csv", ".dat", ".ini", ".conf"]

    def extract(self, file_path: str) -> DocumentContent:
        logger.info(f"Extracting text from {file_path}")
        encoding = getattr(self.config, "encoding", "auto")
        text = ""
        detected_encoding = encoding

        try:
            if encoding == "auto":
                text, detected_encoding = EncodingDetector.detect_and_read(file_path)
            else:
                with open(file_path, "r", encoding=encoding, errors="ignore") as f:
                    text = f.read()
                detected_encoding = encoding
        except Exception as e:
            logger.warning(f"Primary read failed for {file_path}: {e}, trying auto-detection")
            try:
                text, detected_encoding = EncodingDetector.detect_and_read(file_path)
            except Exception as e2:
                logger.error(f"Failed to read {file_path}: {e2}")
                text = ""
                detected_encoding = "failed"

        return DocumentContent(
            text=text,
            metadata={
                "file_name": os.path.basename(file_path),
                "file_size": os.path.getsize(file_path),
                "encoding": detected_encoding,
            },
            tables=[],
            images=[],
            file_path=file_path,
            file_type="text",
        )


class PDFAdapter(BaseFormatAdapter):
    def can_handle(self, file_path: str) -> bool:
        ext = Path(file_path).suffix.lower()
        return ext == ".pdf"

    def extract(self, file_path: str) -> DocumentContent:
        logger.info(f"Extracting text from PDF: {file_path}")
        text = ""
        tables: List[List[List[str]]] = []
        images: List[bytes] = []
        is_encrypted = False
        page_count = 0

        try:
            import PyPDF2
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)

                if reader.is_encrypted:
                    is_encrypted = True
                    logger.warning(f"PDF is encrypted: {file_path}")
                    try:
                        reader.decrypt("")
                    except Exception as e:
                        logger.warning(f"Failed to decrypt PDF with empty password: {e}")

                page_count = len(reader.pages)
                for page_num, page in enumerate(reader.pages):
                    try:
                        page_text = page.extract_text()
                        if page_text:
                            text += page_text + "\n\n"
                    except Exception as e:
                        logger.warning(f"Failed to extract page {page_num}: {e}")
                        continue

        except ImportError:
            logger.warning("PyPDF2 not available, trying pdfplumber as fallback")
            try:
                import pdfplumber
                with pdfplumber.open(file_path) as pdf:
                    page_count = len(pdf.pages)
                    for page in pdf.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text += page_text + "\n\n"
            except ImportError:
                logger.warning("pdfplumber not available, using fallback text extraction")
                text = f"[PDF Content from {os.path.basename(file_path)}]"
            except Exception as e:
                logger.error(f"Failed to extract PDF with pdfplumber: {e}")
                text = f"[PDF Content from {os.path.basename(file_path)}]"
        except Exception as e:
            logger.error(f"Failed to extract PDF {file_path}: {e}")

        text = self._clean_pdf_text(text)

        return DocumentContent(
            text=text,
            metadata={
                "file_name": os.path.basename(file_path),
                "file_size": os.path.getsize(file_path),
                "pages": page_count,
                "is_encrypted": is_encrypted,
            },
            tables=tables,
            images=images,
            file_path=file_path,
            file_type="pdf",
        )

    @staticmethod
    def _clean_pdf_text(text: str) -> str:
        if not text:
            return text
        text = re.sub(r"\x00", "", text)
        text = re.sub(r"[ \t\r\f\v]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        text = re.sub(r"(\w)\s+-\s+(\w)", r"\1-\2", text)
        return text.strip()


class WordAdapter(BaseFormatAdapter):
    def can_handle(self, file_path: str) -> bool:
        ext = Path(file_path).suffix.lower()
        return ext in [".docx", ".doc"]

    def extract(self, file_path: str) -> DocumentContent:
        logger.info(f"Extracting text from Word document: {file_path}")
        text = ""
        tables: List[List[List[str]]] = []
        images: List[bytes] = []

        try:
            from docx import Document
            doc = Document(file_path)
            for para in doc.paragraphs:
                para_text = self._clean_word_text(para.text)
                if para_text:
                    text += para_text + "\n"
            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [self._clean_word_text(cell.text) for cell in row.cells]
                    table_data.append(row_data)
                tables.append(table_data)
        except ImportError:
            logger.warning("python-docx not available, trying alternative extraction")
            text = self._extract_with_ole(file_path)
        except Exception as e:
            logger.warning(f"python-docx extraction failed: {e}, trying alternative")
            text = self._extract_with_ole(file_path)

        if not text:
            text = f"[Word Document: {os.path.basename(file_path)}]"

        return DocumentContent(
            text=text,
            metadata={
                "file_name": os.path.basename(file_path),
                "file_size": os.path.getsize(file_path),
                "tables_count": len(tables),
            },
            tables=tables,
            images=images,
            file_path=file_path,
            file_type="word",
        )

    @staticmethod
    def _clean_word_text(text: str) -> str:
        if not text:
            return ""
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
        text = re.sub(r"\s+", " ", text)
        return text.strip()

    @staticmethod
    def _extract_with_ole(file_path: str) -> str:
        try:
            import subprocess
            result = subprocess.run(
                ["catdoc", "-d", "utf-8", file_path],
                capture_output=True,
                text=True,
                timeout=30
            )
            if result.returncode == 0 and result.stdout:
                return WordAdapter._clean_word_text(result.stdout)
        except Exception as e:
            logger.warning(f"catdoc extraction failed: {e}")
        return ""


class ExcelAdapter(BaseFormatAdapter):
    def can_handle(self, file_path: str) -> bool:
        ext = Path(file_path).suffix.lower()
        return ext in [".xlsx", ".xls", ".csv"]

    def extract(self, file_path: str) -> DocumentContent:
        logger.info(f"Extracting data from Excel/CSV: {file_path}")
        text = ""
        tables: List[List[List[str]]] = []
        images: List[bytes] = []

        ext = Path(file_path).suffix.lower()
        try:
            if ext == ".csv":
                import csv
                csv_text, detected_encoding = EncodingDetector.detect_and_read(file_path)
                if csv_text:
                    import io
                    reader = csv.reader(io.StringIO(csv_text))
                    table_data = [row for row in reader]
                    tables.append(table_data)
                    text = "\n".join([",".join(row) for row in table_data])
            else:
                import openpyxl
                wb = openpyxl.load_workbook(file_path, data_only=True)
                for sheet_name in wb.sheetnames:
                    sheet = wb[sheet_name]
                    table_data = []
                    for row in sheet.iter_rows(values_only=True):
                        row_data = [str(cell) if cell is not None else "" for cell in row]
                        table_data.append(row_data)
                    tables.append(table_data)
                    text += f"Sheet: {sheet_name}\n"
                    text += "\n".join([",".join(row) for row in table_data]) + "\n\n"
        except ImportError:
            logger.warning("openpyxl not available, using fallback")
            text = f"[Spreadsheet: {os.path.basename(file_path)}]"
        except Exception as e:
            logger.error(f"Failed to extract Excel/CSV {file_path}: {e}")

        return DocumentContent(
            text=text,
            metadata={
                "file_name": os.path.basename(file_path),
                "file_size": os.path.getsize(file_path),
                "sheets_count": len(tables),
            },
            tables=tables,
            images=images,
            file_path=file_path,
            file_type="excel",
        )


class HTMLAdapter(BaseFormatAdapter):
    def can_handle(self, file_path: str) -> bool:
        ext = Path(file_path).suffix.lower()
        return ext in [".html", ".htm", ".xml", ".json"]

    def extract(self, file_path: str) -> DocumentContent:
        logger.info(f"Extracting text from HTML/XML/JSON: {file_path}")
        text = ""
        tables: List[List[List[str]]] = []
        images: List[bytes] = []

        ext = Path(file_path).suffix.lower()
        try:
            if ext == ".json":
                import json
                file_content, detected_encoding = EncodingDetector.detect_and_read(file_path)
                if file_content:
                    data = json.loads(file_content)
                    text = json.dumps(data, ensure_ascii=False, indent=2)
            else:
                from bs4 import BeautifulSoup
                file_content, detected_encoding = EncodingDetector.detect_and_read(file_path)
                if file_content:
                    soup = BeautifulSoup(file_content, "html.parser")
                    text = soup.get_text(separator="\n", strip=True)
                    for table in soup.find_all("table"):
                        table_data = []
                        for row in table.find_all("tr"):
                            cells = row.find_all(["th", "td"])
                            row_data = [cell.get_text(strip=True) for cell in cells]
                            table_data.append(row_data)
                        if table_data:
                            tables.append(table_data)
        except ImportError:
            logger.warning("beautifulsoup4 not available, using fallback")
            with open(file_path, "r", encoding=getattr(self.config, "encoding", "utf-8"), errors="ignore") as f:
                text = f.read()
        except Exception as e:
            logger.error(f"Failed to extract HTML/XML/JSON {file_path}: {e}")

        return DocumentContent(
            text=text,
            metadata={
                "file_name": os.path.basename(file_path),
                "file_size": os.path.getsize(file_path),
                "tables_count": len(tables),
            },
            tables=tables,
            images=images,
            file_path=file_path,
            file_type="html",
        )


class EmailAdapter(BaseFormatAdapter):
    def can_handle(self, file_path: str) -> bool:
        ext = Path(file_path).suffix.lower()
        return ext in [".eml", ".msg"]

    def extract(self, file_path: str) -> DocumentContent:
        logger.info(f"Extracting content from email: {file_path}")
        text = ""
        tables: List[List[List[str]]] = []
        images: List[bytes] = []

        try:
            import email
            from email import policy
            with open(file_path, "rb") as f:
                msg = email.message_from_binary_file(f, policy=policy.default)
                subject = msg.get("Subject", "")
                sender = msg.get("From", "")
                recipient = msg.get("To", "")
                date = msg.get("Date", "")

                text_parts = []
                text_parts.append(f"Subject: {subject}")
                text_parts.append(f"From: {sender}")
                text_parts.append(f"To: {recipient}")
                text_parts.append(f"Date: {date}")
                text_parts.append("\nBody:")

                if msg.is_multipart():
                    for part in msg.walk():
                        content_type = part.get_content_type()
                        if content_type == "text/plain":
                            text_parts.append(part.get_content())
                        elif content_type == "text/html":
                            from bs4 import BeautifulSoup
                            soup = BeautifulSoup(part.get_content(), "html.parser")
                            text_parts.append(soup.get_text())
                else:
                    text_parts.append(msg.get_content())

                text = "\n".join(text_parts)
        except ImportError:
            logger.warning("email parsing dependencies not available")
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        except Exception as e:
            logger.error(f"Failed to extract email {file_path}: {e}")

        return DocumentContent(
            text=text,
            metadata={
                "file_name": os.path.basename(file_path),
                "file_size": os.path.getsize(file_path),
            },
            tables=tables,
            images=images,
            file_path=file_path,
            file_type="email",
        )


class FormatAdapterRegistry:
    def __init__(self, config: Any = None):
        self.config = config
        self._adapters: List[BaseFormatAdapter] = []
        self._register_default_adapters()

    def _register_default_adapters(self):
        default_adapters = [
            TextAdapter(self.config),
            PDFAdapter(self.config),
            WordAdapter(self.config),
            ExcelAdapter(self.config),
            HTMLAdapter(self.config),
            EmailAdapter(self.config),
        ]
        for adapter in default_adapters:
            self.register_adapter(adapter)

    def register_adapter(self, adapter: BaseFormatAdapter):
        self._adapters.append(adapter)
        logger.debug(f"Registered adapter: {adapter.__class__.__name__}")

    def get_adapter(self, file_path: str) -> Optional[BaseFormatAdapter]:
        for adapter in self._adapters:
            if adapter.can_handle(file_path):
                return adapter
        return None

    def is_supported(self, file_path: str) -> bool:
        return self.get_adapter(file_path) is not None

    def extract(self, file_path: str) -> Optional[DocumentContent]:
        adapter = self.get_adapter(file_path)
        if adapter:
            return adapter.extract(file_path)
        logger.warning(f"No adapter found for {file_path}")
        return None
