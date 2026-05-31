import os
import json
import logging
import sqlite3
from typing import List, Dict, Any, Optional
from pathlib import Path
from datetime import datetime

from .result_formatter import StructuredResult, ResultFormatter

logger = logging.getLogger(__name__)


class FileStorage:
    def __init__(self, config: Any = None):
        self.config = config
        self.output_dir = getattr(config, "output_dir", "./output") if config else "./output"
        self.save_formats = getattr(config, "save_formats", ["json", "markdown"]) if config else ["json", "markdown"]
        self.enable_versioning = getattr(config, "enable_versioning", True) if config else True
        self.compression = getattr(config, "compression", False) if config else False

        self._ensure_output_dir()

    def _ensure_output_dir(self) -> None:
        Path(self.output_dir).mkdir(parents=True, exist_ok=True)

    def save(self, result: StructuredResult, formats: Optional[List[str]] = None) -> Dict[str, str]:
        save_formats = formats or self.save_formats
        saved_files = {}

        base_name = self._generate_base_name(result)

        if "json" in save_formats:
            file_path = self._save_json(result, base_name)
            saved_files["json"] = file_path

        if "markdown" in save_formats or "md" in save_formats:
            file_path = self._save_markdown(result, base_name)
            saved_files["markdown"] = file_path

        if "excel" in save_formats or "xlsx" in save_formats:
            file_path = self._save_excel(result, base_name)
            if file_path:
                saved_files["excel"] = file_path

        logger.info(f"Saved result {result.document_id} to {len(saved_files)} formats")
        return saved_files

    def save_batch(self, results: List[StructuredResult], formats: Optional[List[str]] = None) -> List[Dict[str, str]]:
        saved_files_list = []
        for result in results:
            saved_files = self.save(result, formats)
            saved_files_list.append(saved_files)
        return saved_files_list

    def _save_json(self, result: StructuredResult, base_name: str) -> str:
        file_name = f"{base_name}.json"
        file_path = os.path.join(self.output_dir, file_name)

        content = ResultFormatter.format_json(result)
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(content)

        return file_path

    def _save_markdown(self, result: StructuredResult, base_name: str) -> str:
        file_name = f"{base_name}.md"
        file_path = os.path.join(self.output_dir, file_name)

        content = ResultFormatter.format_markdown(result)
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(content)

        return file_path

    def _save_excel(self, result: StructuredResult, base_name: str) -> Optional[str]:
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment

            file_name = f"{base_name}.xlsx"
            file_path = os.path.join(self.output_dir, file_name)

            excel_data = ResultFormatter.format_excel(result)

            wb = openpyxl.Workbook()
            wb.remove(wb.active)

            for sheet_name, data in excel_data.items():
                ws = wb.create_sheet(title=sheet_name[:31])
                for row in data:
                    ws.append(row)

                for cell in ws[1]:
                    cell.font = Font(bold=True)
                    cell.alignment = Alignment(horizontal="center")

                for column in ws.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    ws.column_dimensions[column_letter].width = adjusted_width

            wb.save(file_path)
            return file_path

        except ImportError:
            logger.warning("openpyxl not available, skipping Excel save")
        except Exception as e:
            logger.error(f"Failed to save Excel: {e}")

        return None

    def _generate_base_name(self, result: StructuredResult) -> str:
        safe_name = "".join(c for c in result.document_name if c.isalnum() or c in ("-", "_")).rstrip()
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        if self.enable_versioning:
            return f"{safe_name}_{result.document_id}_{timestamp}"
        return f"{safe_name}_{result.document_id}"

    def load(self, file_path: str) -> Optional[StructuredResult]:
        if not os.path.exists(file_path):
            logger.error(f"File not found: {file_path}")
            return None

        ext = Path(file_path).suffix.lower()
        try:
            if ext == ".json":
                with open(file_path, "r", encoding="utf-8") as f:
                    data = json.load(f)
                return StructuredResult(**data)
            else:
                logger.warning(f"Unsupported format for loading: {ext}")
        except Exception as e:
            logger.error(f"Failed to load result: {e}")

        return None

    def list_results(self) -> List[str]:
        if not os.path.exists(self.output_dir):
            return []

        results = []
        for file in os.listdir(self.output_dir):
            if file.endswith(".json"):
                results.append(os.path.join(self.output_dir, file))
        return sorted(results)

    def delete(self, file_path: str) -> bool:
        try:
            if os.path.exists(file_path):
                os.remove(file_path)
                logger.info(f"Deleted {file_path}")
                return True
        except Exception as e:
            logger.error(f"Failed to delete {file_path}: {e}")
        return False


class DatabaseStorage:
    def __init__(self, config: Any = None):
        self.config = config
        self.db_path = getattr(config, "database_path", "./output/results.db") if config else "./output/results.db"
        self._ensure_db_dir()
        self._init_database()

    def _ensure_db_dir(self) -> None:
        db_dir = os.path.dirname(self.db_path)
        if db_dir:
            Path(db_dir).mkdir(parents=True, exist_ok=True)

    def _init_database(self) -> None:
        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.cursor()
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS documents (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    document_id TEXT UNIQUE NOT NULL,
                    document_name TEXT NOT NULL,
                    file_type TEXT NOT NULL,
                    file_path TEXT NOT NULL,
                    summary TEXT,
                    entities TEXT,
                    relations TEXT,
                    tables TEXT,
                    extraction_results TEXT,
                    metadata TEXT,
                    confidence REAL DEFAULT 0.0,
                    created_at TEXT,
                    updated_at TEXT
                )
            """)
            conn.commit()

    def save(self, result: StructuredResult) -> bool:
        try:
            with sqlite3.connect(self.db_path) as conn:
                cursor = conn.cursor()

                entities_json = json.dumps(result.entities, ensure_ascii=False)
                relations_json = json.dumps(result.relations, ensure_ascii=False)
                tables_json = json.dumps(result.tables, ensure_ascii=False)
                extraction_json = json.dumps(result.extraction_results, ensure_ascii=False)
                metadata_json = json.dumps(result.metadata, ensure_ascii=False)

                cursor.execute("""
                    INSERT OR REPLACE INTO documents
                    (document_id, document_name, file_type, file_path, summary,
                     entities, relations, tables, extraction_results, metadata,
                     confidence, created_at, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    result.document_id,
                    result.document_name,
                    result.file_type,
                    result.file_path,
                    result.summary,
                    entities_json,
                    relations_json,
                    tables_json,
                    extraction_json,
                    metadata_json,
                    result.confidence,
                    result.created_at,
                    datetime.now().isoformat(),
                ))
                conn.commit()
                logger.info(f"Saved {result.document_id} to database")
                return True
        except Exception as e:
            logger.error(f"Failed to save to database: {e}")
            return False

    def load(self, document_id: str) -> Optional[StructuredResult]:
        try:
            with sqlite3.connect(self.db_path) as conn:
                cursor = conn.cursor()
                cursor.execute(
                    "SELECT * FROM documents WHERE document_id = ?",
                    (document_id,)
                )
                row = cursor.fetchone()
                if row:
                    columns = [desc[0] for desc in cursor.description]
                    data = dict(zip(columns, row))
                    return StructuredResult(
                        document_id=data["document_id"],
                        document_name=data["document_name"],
                        file_type=data["file_type"],
                        file_path=data["file_path"],
                        summary=data["summary"] or "",
                        entities=json.loads(data["entities"]) if data["entities"] else [],
                        relations=json.loads(data["relations"]) if data["relations"] else [],
                        tables=json.loads(data["tables"]) if data["tables"] else [],
                        extraction_results=json.loads(data["extraction_results"]) if data["extraction_results"] else {},
                        metadata=json.loads(data["metadata"]) if data["metadata"] else {},
                        confidence=data["confidence"] or 0.0,
                        created_at=data["created_at"],
                    )
        except Exception as e:
            logger.error(f"Failed to load from database: {e}")
        return None

    def list_documents(self, limit: int = 100, offset: int = 0) -> List[Dict[str, Any]]:
        try:
            with sqlite3.connect(self.db_path) as conn:
                conn.row_factory = sqlite3.Row
                cursor = conn.cursor()
                cursor.execute(
                    "SELECT document_id, document_name, file_type, confidence, created_at FROM documents ORDER BY created_at DESC LIMIT ? OFFSET ?",
                    (limit, offset)
                )
                return [dict(row) for row in cursor.fetchall()]
        except Exception as e:
            logger.error(f"Failed to list documents: {e}")
        return []

    def delete(self, document_id: str) -> bool:
        try:
            with sqlite3.connect(self.db_path) as conn:
                cursor = conn.cursor()
                cursor.execute(
                    "DELETE FROM documents WHERE document_id = ?",
                    (document_id,)
                )
                conn.commit()
                return cursor.rowcount > 0
        except Exception as e:
            logger.error(f"Failed to delete from database: {e}")
        return False

    def search(self, keyword: str) -> List[Dict[str, Any]]:
        try:
            with sqlite3.connect(self.db_path) as conn:
                conn.row_factory = sqlite3.Row
                cursor = conn.cursor()
                cursor.execute("""
                    SELECT document_id, document_name, file_type, confidence, created_at
                    FROM documents
                    WHERE document_name LIKE ? OR summary LIKE ?
                    ORDER BY created_at DESC
                """, (f"%{keyword}%", f"%{keyword}%"))
                return [dict(row) for row in cursor.fetchall()]
        except Exception as e:
            logger.error(f"Failed to search: {e}")
        return []


class ResultStorage:
    def __init__(self, config: Any = None):
        self.config = config
        self.file_storage = FileStorage(config)
        self.db_storage = DatabaseStorage(config)

    def save(self, result: StructuredResult, formats: Optional[List[str]] = None) -> Dict[str, Any]:
        file_paths = self.file_storage.save(result, formats)
        db_saved = self.db_storage.save(result)

        return {
            "files": file_paths,
            "database": db_saved,
            "document_id": result.document_id,
        }

    def save_batch(self, results: List[StructuredResult], formats: Optional[List[str]] = None) -> List[Dict[str, Any]]:
        return [self.save(result, formats) for result in results]

    def load(self, identifier: str) -> Optional[StructuredResult]:
        if identifier.endswith(".json"):
            return self.file_storage.load(identifier)
        return self.db_storage.load(identifier)

    def list_files(self) -> List[str]:
        return self.file_storage.list_results()

    def list_documents(self, limit: int = 100, offset: int = 0) -> List[Dict[str, Any]]:
        return self.db_storage.list_documents(limit, offset)

    def search(self, keyword: str) -> List[Dict[str, Any]]:
        return self.db_storage.search(keyword)

    def delete(self, document_id: str) -> bool:
        self.db_storage.delete(document_id)
        for file_path in self.file_storage.list_results():
            if document_id in file_path:
                self.file_storage.delete(file_path)
        return True

    def get_output_dir(self) -> str:
        return self.file_storage.output_dir
